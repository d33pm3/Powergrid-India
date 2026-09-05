from pathlib import Path

from openpyxl import load_workbook

from powergrid_india.parser import parse_master_table
from powergrid_india.reconcile import reconcile
from powergrid_india.workbook import build_workbook_data, create_workbook


EXPECTED_SHEETS = [
    "All_Substations",
    "Regional_Summary",
    "State_Summary",
    "Western_Region",
    "Northern_Region",
    "Southern_Region",
    "Eastern_Region",
    "NorthEastern_Region",
    "Top_50_Substations",
    "Analytics_Dashboard",
]


def test_build_workbook_data_is_internally_consistent():
    recs = reconcile(parse_master_table(), [])
    ds = build_workbook_data(recs)
    assert len(ds) == 9
    assert len(ds["Regional_Summary"]["data"]) == 5
    assert len(ds["State_Summary"]["data"]) == 30
    assert len(ds["Top_50_Substations"]["data"]) == 50
    total = sum(r["Total_Capacity_MW"] for r in ds["All_Substations"]["data"])
    assert total == sum(x["Total_MW"] for x in ds["Regional_Summary"]["data"])
    assert total == 130_500


def test_create_workbook_writes_ten_sheets(tmp_path: Path):
    recs = reconcile(parse_master_table(), [])
    out = tmp_path / "grid.xlsx"
    path, sheets = create_workbook(master_records=recs, output_path=str(out))
    assert Path(path).is_file()
    assert sheets == EXPECTED_SHEETS
    wb = load_workbook(path)
    assert wb.sheetnames == EXPECTED_SHEETS
    ws = wb["All_Substations"]
    assert ws["A1"].value == "Region"
    assert ws["K1"].value == "Data_Source"
    assert ws.max_row == 133  # header + 132
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref
    dash = wb["Analytics_Dashboard"]
    joined = " ".join(str(c.value or "") for row in dash.iter_rows(max_col=3) for c in row)
    assert "BASELINE EPISTEMIC STATUS" in joined
    assert "planning-grade" in joined
    found = {}
    for row in dash.iter_rows(min_col=1, max_col=2, values_only=True):
        if row[0] in {"765/400 kV", "400/220 kV", "220 kV", "132 kV"}:
            found[row[0]] = row[1]
    assert found["765/400 kV"] == 21
    assert found["400/220 kV"] == 104
    assert found["220 kV"] == 6
    assert found["132 kV"] == 1
    assert "Rangpo PS" not in joined
    assert "Khavda BESS PS" in joined
    assert "Pinnapuram PSP PS" in joined
