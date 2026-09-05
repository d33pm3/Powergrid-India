"""10-sheet openpyxl workbook builder (9 data sheets + Analytics Dashboard)."""

from __future__ import annotations

import json
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Mapping

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .reconcile import SCHEMA_FIELDS, coerce_mw

COLUMNS = list(SCHEMA_FIELDS)
REGIONS = [
    ("WR", "Western_Region"),
    ("NR", "Northern_Region"),
    ("SR", "Southern_Region"),
    ("ER", "Eastern_Region"),
    ("NER", "NorthEastern_Region"),
]

H_GREEN = "2E7D32"
H_BLUE = "1C3A6B"
H_TEAL = "00695C"
H_AMBER = "E65100"
H_GREY = "F5F5F5"
TEXT_WH = "FFFFFF"
TEXT_BK = "000000"

STORAGE_TYPES = (
    "BESS",
    "Pump Storage",
    "PSP",
    "Pump Storage/Hydro",
    "Solar/BESS",
    "Solar/Wind/Hybrid/BESS",
    "Solar/Pump Storage",
)

EPISTEMIC_CAVEAT = (
    "Baseline is planning-grade, compiled from public reporting — not a verified "
    "regulatory filing. Rows tagged [SOURCE_FALLBACK:BASELINE] were not confirmed "
    "against a live source this session. Capacity, MVA, and developer-allocation "
    "figures are indicative."
)


def build_workbook_data(master_records: Iterable[Mapping[str, Any]]) -> dict[str, dict[str, Any]]:
    recs = sorted(
        (dict(r) for r in master_records),
        key=lambda r: -coerce_mw(r.get("Total_Capacity_MW", 0)),
    )
    for r in recs:
        r["Total_Capacity_MW"] = coerce_mw(r.get("Total_Capacity_MW"))

    ds: dict[str, dict[str, Any]] = {"All_Substations": {"columns": COLUMNS, "data": recs}}

    reg_agg: dict[str, dict[str, int]] = defaultdict(lambda: {"n": 0, "mw": 0})
    st_agg: dict[str, dict[str, int]] = defaultdict(lambda: {"n": 0, "mw": 0})
    total_mw = 0
    for r in recs:
        mw = coerce_mw(r.get("Total_Capacity_MW", 0))
        reg_agg[str(r.get("Region", "?"))]["n"] += 1
        reg_agg[str(r.get("Region", "?"))]["mw"] += mw
        st_agg[str(r.get("State", "?"))]["n"] += 1
        st_agg[str(r.get("State", "?"))]["mw"] += mw
        total_mw += mw

    ds["Regional_Summary"] = {
        "columns": ["Region", "Substations", "Total_MW", "Pct_Share"],
        "data": [
            {
                "Region": code,
                "Substations": reg_agg[code]["n"],
                "Total_MW": reg_agg[code]["mw"],
                "Pct_Share": f"{reg_agg[code]['mw'] / total_mw * 100:.1f}%" if total_mw else "0%",
            }
            for code, _ in REGIONS
        ],
    }
    ds["State_Summary"] = {
        "columns": ["State", "Substations", "Total_MW"],
        "data": [
            {"State": s, "Substations": v["n"], "Total_MW": v["mw"]}
            for s, v in sorted(st_agg.items(), key=lambda x: -x[1]["mw"])
        ],
    }
    for code, sheet in REGIONS:
        ds[sheet] = {"columns": COLUMNS, "data": [r for r in recs if r.get("Region") == code]}
    ds["Top_50_Substations"] = {"columns": COLUMNS, "data": recs[:50]}

    assert len(ds) == 9, f"expected 9 datasets, built {len(ds)}"
    assert len(ds["Regional_Summary"]["data"]) == 5
    assert sum(coerce_mw(r["Total_Capacity_MW"]) for r in recs) == sum(
        x["Total_MW"] for x in ds["Regional_Summary"]["data"]
    ), "summary drift"
    return ds


def _thin() -> Border:
    s = Side(style="thin")
    return Border(left=s, right=s, top=s, bottom=s)


def _hdr(ws, row, col, val, fill_hex=H_GREEN, sz=11):
    c = ws.cell(row=row, column=col, value=val)
    c.font = Font(bold=True, color=TEXT_WH, size=sz)
    c.fill = PatternFill(start_color=fill_hex, end_color=fill_hex, fill_type="solid")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = _thin()
    return c


def _cell(ws, row, col, val, bold=False, fill_hex=None, align="left"):
    c = ws.cell(row=row, column=col, value=val)
    c.font = Font(bold=bold, color=TEXT_BK, size=10)
    if fill_hex:
        c.fill = PatternFill(start_color=fill_hex, end_color=fill_hex, fill_type="solid")
    c.alignment = Alignment(vertical="center", horizontal=align, wrap_text=True)
    c.border = _thin()
    return c


def _autofit(ws, columns, max_w=55):
    for col_idx, col_name in enumerate(columns, 1):
        max_len = len(str(col_name))
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, min(len(str(cell.value)), max_w))
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 4


def write_data_sheet(wb, sheet_name, sheet_data, idx, fill_hex=H_GREEN):
    ws = wb.active if idx == 0 else wb.create_sheet(sheet_name)
    if idx == 0:
        ws.title = sheet_name
    columns = sheet_data["columns"]
    for ci, cn in enumerate(columns, 1):
        _hdr(ws, 1, ci, cn, fill_hex=fill_hex)
    for ri, rec in enumerate(sheet_data["data"], 2):
        for ci, cn in enumerate(columns, 1):
            val = rec.get(cn, "")
            bg = H_GREY if ri % 2 == 0 else None
            _cell(
                ws,
                ri,
                ci,
                val if val != "" else None,
                fill_hex=bg,
                align="right" if cn in {"Total_Capacity_MW", "Total_MW", "Substations"} else "left",
            )
    _autofit(ws, columns)
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"


def _voltage_tier(voltage: str) -> str | None:
    """Exclusive tier match. Avoids '220 kV' matching inside '400/220 kV'."""
    v = " ".join(str(voltage or "").split())
    if v.startswith("765"):
        return "765/400 kV"
    if v.startswith("400"):
        return "400/220 kV"
    if v.startswith("220"):
        return "220 kV"
    if v.startswith("132"):
        return "132 kV"
    return None


def _is_storage(project_type: str) -> bool:
    pt = str(project_type or "")
    tokens = ("BESS", "Pump Storage", "PSP")
    return any(t in pt for t in tokens)


def write_dashboard(wb, master_records, source_log):
    ws = wb.create_sheet("Analytics_Dashboard")
    row = 1
    records = list(master_records)
    total_mw = sum(coerce_mw(r.get("Total_Capacity_MW", 0)) for r in records)

    _hdr(ws, row, 1, "EXECUTIVE KPI SUMMARY — PAN INDIA RE SUBSTATIONS", fill_hex=H_BLUE, sz=12)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    row += 1
    kpis = [
        ("Total Substations (Post-Reconciliation)", len(records)),
        ("Total PAN India RE Capacity (MW)", total_mw),
        ("Active Substations", sum(1 for r in records if str(r.get("Remarks", "")).startswith("Active"))),
        ("Under Implementation", sum(1 for r in records if str(r.get("Remarks", "")).startswith("Under"))),
        ("Planned Substations", sum(1 for r in records if str(r.get("Remarks", "")).startswith("Planned"))),
        ("GIS Substations", sum(1 for r in records if r.get("Type_AIS_GIS") == "GIS")),
        ("AIS Substations", sum(1 for r in records if r.get("Type_AIS_GIS") == "AIS")),
        (
            "Live-Updated Records",
            sum(1 for r in records if "[SOURCE_LIVE]" in str(r.get("Data_Source", ""))),
        ),
    ]
    for label, val in kpis:
        _hdr(ws, row, 1, label, fill_hex=H_TEAL)
        _cell(ws, row, 2, val, bold=True, align="right")
        row += 1
    row += 1

    _hdr(ws, row, 1, "REGIONAL CAPACITY BREAKDOWN", fill_hex=H_BLUE, sz=12)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    row += 1
    for h, c in [("Region", 1), ("Substations", 2), ("Total MW", 3), ("% of PAN India", 4)]:
        _hdr(ws, row, c, h, fill_hex=H_TEAL)
    row += 1
    for region, name in [
        ("WR", "Western"),
        ("SR", "Southern"),
        ("NR", "Northern"),
        ("ER", "Eastern"),
        ("NER", "North-Eastern"),
    ]:
        recs = [r for r in records if r.get("Region") == region]
        mw = sum(coerce_mw(r.get("Total_Capacity_MW", 0)) for r in recs)
        pct = f"{mw / total_mw * 100:.1f}%" if total_mw else "0%"
        _cell(ws, row, 1, f"{region} ({name})")
        _cell(ws, row, 2, len(recs), align="right")
        _cell(ws, row, 3, mw, align="right")
        _cell(ws, row, 4, pct, align="right")
        row += 1
    row += 1

    _hdr(ws, row, 1, "VOLTAGE TIER ANALYSIS", fill_hex=H_BLUE, sz=12)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    row += 1
    for h, c in [("Voltage Level", 1), ("Substations", 2), ("Total MW", 3), ("Tier Classification", 4)]:
        _hdr(ws, row, c, h, fill_hex=H_TEAL)
    row += 1
    tiers = [
        ("765/400 kV", "Tier-1 CTU Backbone — Bulk HVAC/HVDC evacuation"),
        ("400/220 kV", "Tier-2 CTU Distribution — Main RE pooling nodes"),
        ("220 kV", "State-level / Special Purpose (PSP, frontier)"),
        ("132 kV", "Sub-state / Frontier / Small capacity"),
    ]
    for vl, label in tiers:
        recs = [r for r in records if _voltage_tier(r.get("Voltage_Level", "")) == vl]
        mw = sum(coerce_mw(r.get("Total_Capacity_MW", 0)) for r in recs)
        _cell(ws, row, 1, vl)
        _cell(ws, row, 2, len(recs), align="right")
        _cell(ws, row, 3, mw, align="right")
        _cell(ws, row, 4, label)
        row += 1
    row += 1

    _hdr(ws, row, 1, "PROJECT TYPE MIX", fill_hex=H_BLUE, sz=12)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    row += 1
    for h, c in [("Project Type", 1), ("Substations", 2), ("Total MW", 3)]:
        _hdr(ws, row, c, h, fill_hex=H_TEAL)
    row += 1
    type_map: dict[str, dict[str, int]] = defaultdict(lambda: {"count": 0, "mw": 0})
    for r in records:
        pt = str(r.get("Project_Type") or "Unknown")
        type_map[pt]["count"] += 1
        type_map[pt]["mw"] += coerce_mw(r.get("Total_Capacity_MW", 0))
    for pt, vals in sorted(type_map.items(), key=lambda x: -x[1]["mw"]):
        _cell(ws, row, 1, pt)
        _cell(ws, row, 2, vals["count"], align="right")
        _cell(ws, row, 3, vals["mw"], align="right")
        row += 1
    row += 1

    _hdr(ws, row, 1, "BESS & PUMPED STORAGE NODES", fill_hex=H_BLUE, sz=12)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    row += 1
    for h, c in [("Substation", 1), ("State", 2), ("Region", 3), ("MW", 4), ("Status", 5)]:
        _hdr(ws, row, c, h, fill_hex=H_TEAL)
    row += 1
    storage_recs = [r for r in records if _is_storage(r.get("Project_Type", ""))]
    for r in sorted(storage_recs, key=lambda x: -coerce_mw(x.get("Total_Capacity_MW", 0))):
        status = str(r.get("Remarks", "")).split(" - ")[0]
        fill = H_AMBER if status == "Planned" else None
        _cell(ws, row, 1, r.get("Substation_Name"), fill_hex=fill)
        _cell(ws, row, 2, r.get("State"), fill_hex=fill)
        _cell(ws, row, 3, r.get("Region"), fill_hex=fill)
        _cell(ws, row, 4, coerce_mw(r.get("Total_Capacity_MW")), align="right", fill_hex=fill)
        _cell(ws, row, 5, status, fill_hex=fill)
        row += 1
    row += 1

    _hdr(ws, row, 1, "GREEN ENERGY CORRIDOR (GEC) SUBSTATIONS", fill_hex=H_BLUE, sz=12)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    row += 1
    for h, c in [("GEC Phase", 1), ("Substation", 2), ("State", 3), ("MW", 4), ("Status", 5)]:
        _hdr(ws, row, c, h, fill_hex=H_TEAL)
    row += 1
    gec_recs = [r for r in records if "GEC" in str(r.get("Remarks", ""))]
    for r in sorted(gec_recs, key=lambda x: -coerce_mw(x.get("Total_Capacity_MW", 0))):
        rmk = str(r.get("Remarks", ""))
        phase = (
            "GEC Phase-I"
            if "Phase-I" in rmk and "Phase-II" not in rmk
            else "GEC Phase-II"
            if "Phase-II" in rmk
            else "GEC (Phase unspecified)"
        )
        status = rmk.split(" - ")[0]
        _cell(ws, row, 1, phase)
        _cell(ws, row, 2, r.get("Substation_Name"))
        _cell(ws, row, 3, r.get("State"))
        _cell(ws, row, 4, coerce_mw(r.get("Total_Capacity_MW")), align="right")
        _cell(ws, row, 5, status)
        row += 1
    row += 1

    _hdr(ws, row, 1, "TOP 15 SUBSTATIONS BY CAPACITY (MW)", fill_hex=H_BLUE, sz=12)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    row += 1
    for h, c in [("Rank", 1), ("Substation", 2), ("State", 3), ("Region", 4), ("MW", 5), ("Status", 6)]:
        _hdr(ws, row, c, h, fill_hex=H_TEAL)
    row += 1
    ranked = sorted(records, key=lambda x: -coerce_mw(x.get("Total_Capacity_MW", 0)))[:15]
    for rank, r in enumerate(ranked, start=1):
        status = str(r.get("Remarks", "")).split(" - ")[0]
        _cell(ws, row, 1, rank, align="center")
        _cell(ws, row, 2, r.get("Substation_Name"))
        _cell(ws, row, 3, r.get("State"))
        _cell(ws, row, 4, r.get("Region"))
        _cell(ws, row, 5, coerce_mw(r.get("Total_Capacity_MW")), align="right")
        _cell(ws, row, 6, status)
        row += 1
    row += 1

    _hdr(ws, row, 1, "DATA SOURCE & LIVE RESEARCH LOG", fill_hex=H_BLUE, sz=12)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    row += 1
    for h, c in [("Source", 1), ("URL", 2), ("Status", 3), ("Records Updated", 4), ("Timestamp", 5)]:
        _hdr(ws, row, c, h, fill_hex=H_TEAL)
    row += 1
    ts = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    for src in source_log or []:
        _cell(ws, row, 1, src.get("name"))
        _cell(ws, row, 2, src.get("url"))
        _cell(ws, row, 3, src.get("status"))
        _cell(ws, row, 4, src.get("records_updated", 0), align="right")
        _cell(ws, row, 5, src.get("timestamp", ts))
        row += 1
    row += 1
    _hdr(ws, row, 1, "BASELINE EPISTEMIC STATUS", fill_hex=H_BLUE, sz=12)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    row += 1
    _cell(ws, row, 1, EPISTEMIC_CAVEAT)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)

    for col in range(1, 7):
        max_len = 15
        for r2 in ws.iter_rows(min_col=col, max_col=col):
            for cell in r2:
                if cell.value:
                    max_len = max(max_len, min(len(str(cell.value)), 60))
        ws.column_dimensions[get_column_letter(col)].width = max_len + 3
    ws.auto_filter.ref = "A1:F1"
    ws.freeze_panes = "A2"


def default_source_log() -> list[dict[str, Any]]:
    return [
        {"name": "CTUIL", "url": "https://www.ctuil.in", "status": "[SOURCE_FALLBACK]", "records_updated": 0},
        {"name": "PGCIL", "url": "https://www.powergrid.in", "status": "[SOURCE_FALLBACK]", "records_updated": 0},
        {"name": "CEA", "url": "https://www.cea.nic.in", "status": "[SOURCE_FALLBACK]", "records_updated": 0},
        {"name": "MNRE", "url": "https://www.mnre.gov.in", "status": "[SOURCE_FALLBACK]", "records_updated": 0},
        {
            "name": "GEC Docs",
            "url": "https://www.powergrid.in/project/green-energy-corridor",
            "status": "[SOURCE_FALLBACK]",
            "records_updated": 0,
        },
        {
            "name": "Baseline",
            "url": "Skill powergrid-india v2.1.2 (planning-grade, indicative)",
            "status": "ALWAYS_USED",
            "records_updated": 132,
        },
    ]


def create_workbook(
    master_records=None,
    source_log=None,
    output_path="outputs/PAN_India_RE_Substations_Master_Data.xlsx",
    json_path=None,
):
    """Records-first builder. json_path is a legacy fallback only."""
    if master_records is not None:
        workbook_data = build_workbook_data(master_records)
    elif json_path:
        with open(json_path, "r", encoding="utf-8") as f:
            payload = json.load(f)
        workbook_data = payload["workbook_data"] if "workbook_data" in payload else payload
        master_records = workbook_data["All_Substations"]["data"]
    else:
        raise ValueError("Provide master_records (preferred) or json_path")

    out = Path(output_path)
    if out.parent and str(out.parent) not in {"", "."}:
        out.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    sheet_order = [
        "All_Substations",
        "Regional_Summary",
        "State_Summary",
        "Western_Region",
        "Northern_Region",
        "Southern_Region",
        "Eastern_Region",
        "NorthEastern_Region",
        "Top_50_Substations",
    ]
    for idx, sname in enumerate(sheet_order):
        if sname not in workbook_data:
            continue
        write_data_sheet(wb, sname, workbook_data[sname], idx)

    write_dashboard(wb, master_records, source_log if source_log is not None else default_source_log())
    wb.save(out)
    return str(out), list(wb.sheetnames)
